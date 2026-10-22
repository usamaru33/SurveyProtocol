#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""全文評価の閲覧アプリ(review_<id>.html)が書き出した進捗 JSON を、判定シート(xlsx)へ反映する。

書くのは入力列だけ(6 基準・根拠・メモ・事実の抽出)。「判定(自動)」と、自動のままの「除外理由」は
式のまま残す。除外理由は、アプリで**選び直した**とき(自動と違うとき)だけ値として書く。

検証してから書く。止まる条件:
  - 進捗 JSON の形式・シート ID が違う(他の評価者の JSON を自分のシートへ反映してはいけない)
  - 担当外の record_id が混じっている
  - 基準が Y/N/? でない / 除外理由・語彙の列が語彙の外 / n_total が整数でも「記載なし」でもない
  - 既にシートに別の値が入っているセルを上書きしようとしている(--force で許可)
書く前に、元の xlsx を `<id>/backup/` へ退避する。Excel でシートを開いたままだと書けない(閉じてから実行)。

使い方:
    python -X utf8 scripts/apply_fulltext_review.py --id author --input screening/phase4/author/progress_author.json
    python -X utf8 scripts/apply_fulltext_review.py --id author --input ... --dry-run
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

_HERE = Path(__file__).resolve().parent
ROOT = _HERE.parent
sys.path.insert(0, str(_HERE))

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("[ERROR] openpyxl が必要です:  pip install openpyxl")

from make_fulltext_sheets import (  # noqa: E402
    COLS, COL_INDEX, CRIT_KEYS, EXTRACT, REASON_VALUES, YNQ, auto_reason, expected_decision,
)
from make_screening_sheets import REVIEWERS  # noqa: E402

PHASE4 = ROOT / "screening" / "phase4"
EX = {e[0]: e for e in EXTRACT}


def norm(v) -> str:
    return "" if v is None else str(v).strip()


def validate(entry: dict, tag: str) -> tuple[dict, list[str]]:
    """1 件分の JSON を、シートに書く {列キー: 値} と、エラー一覧にする。"""
    err: list[str] = []
    out: dict[str, str] = {}
    c = entry.get("c") or {}
    for k in CRIT_KEYS:
        v = norm(c.get(k))
        if v and v not in YNQ:
            err.append(f"{tag} {k} が Y/N/? でない: 「{v}」")
        out[k] = v
    # 除外理由は、自動と違う(選び直した)ときだけ値として書く
    reason = norm(entry.get("reason"))
    if reason and reason not in REASON_VALUES:
        err.append(f"{tag} 除外理由が語彙外: 「{reason}」")
    crit = {k: out[k] for k in CRIT_KEYS}
    out["reason"] = reason if (reason and expected_decision(crit) == "Exclude" and reason != auto_reason(crit)) else ""
    out["evidence"], out["memo"] = norm(entry.get("evidence")), norm(entry.get("memo"))
    for k, v in (entry.get("ext") or {}).items():
        if k not in EX:
            err.append(f"{tag} 未知の抽出列: {k}")
            continue
        _k, _label, _req, kind, vocab, *_ = EX[k]
        if kind == "multi":
            toks = [norm(t) for t in (v if isinstance(v, list) else str(v).split(";")) if norm(t)]
            bad = [t for t in toks if t not in vocab]
            if bad:
                err.append(f"{tag} {k} が語彙外: {bad}")
            out[k] = "; ".join(toks)
        else:
            s = norm(v)
            if kind in ("vocab", "ynq") and s and s not in vocab:
                err.append(f"{tag} {k} が語彙外: 「{s}」")
            if kind == "int" and s and not (s.isdigit() or s == "記載なし"):
                err.append(f"{tag} {k} が整数でも「記載なし」でもない: 「{s}」")
            out[k] = int(s) if kind == "int" and s.isdigit() else s
    return out, err


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--id", required=True, choices=list(REVIEWERS))
    ap.add_argument("--input", required=True, type=Path, help="進捗 JSON(アプリの S キーで書き出したもの)")
    ap.add_argument("--dir", type=Path, default=PHASE4)
    ap.add_argument("--dry-run", action="store_true", help="検証だけして書かない")
    ap.add_argument("--force", action="store_true", help="シートに既にある別の値を上書きする")
    args = ap.parse_args()

    data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict) or data.get("format") != "phase4-progress":
        sys.exit(f"[ERROR] {args.input.name} は全文評価アプリの進捗 JSON でない (format={data.get('format') if isinstance(data, dict) else '?'})")
    if int(data.get("version") or 0) > 1:
        sys.exit("[ERROR] 進捗 JSON の version がこのスクリプトより新しい。scripts/apply_fulltext_review.py を更新すること")
    if data.get("sheet") != args.id:
        sys.exit(f"[ERROR] シート ID が違う(ファイル: {data.get('sheet')!r} / --id: {args.id!r})。他人の判定を自分のシートへ反映してはいけない")
    entries = data.get("entries") or {}
    print(f"[INFO] 進捗 JSON {args.input.name}: 保存 {data.get('saved_at')} / {len(entries)} 件 / fingerprint {data.get('fingerprint')}")

    path = args.dir / args.id / f"phase4_sheet_{args.id}.xlsx"
    if not path.exists():
        sys.exit(f"[ERROR] {path} が無い")
    wb = load_workbook(path)
    ws = wb["判定"]
    if [c.value for c in ws[1]] != [h for _k, h, _w, _t in COLS]:
        sys.exit("[ERROR] シートの列がこのスクリプトの想定と違う(シートを作り直したか、スクリプトが古い)")
    rows = {ws.cell(r, COL_INDEX["record_id"]).value: r for r in range(2, ws.max_row + 1)}

    errors: list[str] = []
    conflicts: list[str] = []
    writes: list[tuple[int, str, object]] = []
    for rid, entry in entries.items():
        tag = f"{rid}"
        if rid not in rows:
            errors.append(f"{tag} 担当外(シートに無い)の record_id")
            continue
        vals, err = validate(entry, tag)
        errors += err
        r = rows[rid]
        for k, v in vals.items():
            cell = ws.cell(r, COL_INDEX[k])
            old = cell.value
            if k == "reason" and v == "":
                if old and not str(old).startswith("="):
                    print(f"[WARN] {rid} 除外理由: シートに選び直した値「{old}」があるが、アプリは自動。シートの値を残す")
                continue                                   # 自動のまま(式を残す)
            if v == "" and (old is None or old == ""):
                continue
            if norm(old) not in ("", norm(v)) and not (isinstance(old, str) and old.startswith("=")):
                conflicts.append(f"{tag} {k}: シート「{norm(old)}」 ≠ アプリ「{norm(v)}」")
            writes.append((r, k, v))
    if errors:
        print("[ERROR] 検証で止まった(何も書いていない):")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    if conflicts and not args.force:
        print(f"[ERROR] シートに既に別の値が入っているセルが {len(conflicts)} 件ある(何も書いていない)。")
        print("  Excel で直接書いた内容かもしれない。確認して、アプリの値で上書きしてよければ --force:")
        for c in conflicts[:20]:
            print("  -", c)
        sys.exit(1)

    n = len({r for r, _k, _v in writes})
    print(f"[INFO] 反映する: {n} 件・{len(writes)} セル" + (f"(上書き {len(conflicts)} セルを含む)" if conflicts else ""))
    if args.dry_run:
        print("[DRY-RUN] 書かずに終わる")
        return
    bdir = args.dir / args.id / "backup"
    bdir.mkdir(exist_ok=True)
    bak = bdir / f"phase4_sheet_{args.id}_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    shutil.copyfile(path, bak)
    for r, k, v in writes:
        ws.cell(r, COL_INDEX[k]).value = v if v != "" else None
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"[ERROR] {path.name} に書けない。Excel で開いたままなら閉じてから、もう一度実行すること(何も書き換わっていない)")
    print(f"[OK] {path} を更新した。元のファイルは {bak} に退避した。")


if __name__ == "__main__":
    main()
