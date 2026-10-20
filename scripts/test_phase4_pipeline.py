#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_phase4_pipeline.py — Phase 4 のシート生成と集計を、合成データで端から端まで検証する
================================================================================
本物の判定シートには触れない(一時フォルダに生成する)。次を確認する:
  1. 割当が均等化 案 B(50/50/49、評価回数 100/99/99)で、決定論的である
  2. 生成したシートに、注入した記入ミス(根拠なしの Exclude・語彙外の除外理由)を score_fulltext.check_sheet が
     検出する。判定と除外理由は式で自動に出る(Excel で保存していないファイルでも 6 列から計算される)
  3. 合成の判定から κ・協議リスト・最終判定・抽出表が書き出せる
実行:  python -X utf8 scripts/test_phase4_pipeline.py
"""
import csv
import random
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from make_fulltext_sheets import COL_INDEX, CRIT_KEYS, assign  # noqa: E402
import score_fulltext as sf  # noqa: E402

PY = [sys.executable, "-X", "utf8"]
REASON = {"P": "P: 対象者が不適合", "I_hmd": "I: HMD-VR でない", "I_scale": "I: スケール操作/多感覚刺激が無い",
          "O": "O: スケール知覚の測定が無い", "S": "S: ユーザー実験が無い", "C": "その他"}


def fill(base: Path) -> None:
    for who in ("author", "kataoka", "watanabe"):
        p = base / who / f"phase4_sheet_{who}.xlsx"
        wb = openpyxl.load_workbook(p)
        ws = wb["判定"]
        for r in range(2, ws.max_row + 1):
            rid = ws.cell(r, 1).value
            truth, rnd = random.Random(rid).random(), random.Random(rid + who)
            vals = {}
            for k in CRIT_KEYS:
                v = "N" if (truth < 0.35 and k in ("I_scale", "O")) else "Y"
                if rnd.random() < 0.08:
                    v = "?"
                vals[k] = v
            vs = list(vals.values())
            d = "Exclude" if "N" in vs else ("Unsure" if "?" in vs else "Include")
            for k in CRIT_KEYS:
                ws.cell(r, COL_INDEX[k]).value = vals[k]
            # 判定と除外理由は式(自動)なので書かない。Exclude の根拠だけ書く
            if d == "Exclude":
                ws.cell(r, COL_INDEX["evidence"]).value = 'p.3 "..."'
            if ws.cell(r, COL_INDEX["extractor"]).value == "Y" and d in ("Include", "Unsure"):
                for k, v in {"n_total": "24", "population": "大学生", "device": "HTC Vive", "scale_target": "身体全体",
                             "scale_factor": "0.3x", "non_visual": "なし", "task": "距離推定; 歩行・移動",
                             "outcome_measures": "距離推定(行動)", "main_finding": "結果",
                             "attribution_quote": "記載なし"}.items():
                    ws.cell(r, COL_INDEX[k]).value = v
        wb.save(p)


def main() -> None:
    ids = [f"R{i:09x}" for i in range(149)]
    plan = assign(ids)
    pairs = [sum(1 for p in plan.values() if p["pair"] == k) for k in ("author×kataoka", "author×watanabe", "kataoka×watanabe")]
    assert pairs == [50, 50, 49], pairs
    assert assign(ids) == assign(list(reversed(ids))), "割当が入力順に依存している"
    ev = {w: sum(w in (p["reviewer_a"], p["reviewer_b"]) for p in plan.values()) for w in ("author", "kataoka", "watanabe")}
    assert ev == {"author": 100, "kataoka": 99, "watanabe": 99}, ev
    print("OK 割当: 50/50/49、評価回数 100/99/99、決定論的")

    with tempfile.TemporaryDirectory() as td:
        base = Path(td) / "phase4"
        subprocess.run(PY + [str(HERE / "make_fulltext_sheets.py"), "--out", str(base)], check=True, capture_output=True)
        fill(base)
        # 記入ミスの注入(著者のシート)
        p = base / "author" / "phase4_sheet_author.xlsx"
        wb = openpyxl.load_workbook(p)
        ws = wb["判定"]
        for k, v in (("P", "Y"), ("I_hmd", "Y"), ("I_scale", "N"), ("C", "Y"), ("O", "Y"), ("S", "Y")):
            ws.cell(2, COL_INDEX[k]).value = v
        ws.cell(2, COL_INDEX["evidence"]).value = None               # Exclude なのに根拠なし
        for k, v in (("P", "Y"), ("I_hmd", "Y"), ("I_scale", "Y"), ("C", "Y"), ("O", "N"), ("S", "Y")):
            ws.cell(4, COL_INDEX[k]).value = v
        ws.cell(4, COL_INDEX["reason"]).value = "患者対象"             # 語彙外(式を手入力で上書き)
        ws.cell(4, COL_INDEX["evidence"]).value = "p.1"
        wb.save(p)
        sheets = {w: sf.load_sheet(base / w / f"phase4_sheet_{w}.xlsx") for w in ("author", "kataoka", "watanabe")}
        assigned = {a["record_id"]: a for a in sf.load_assignment(base)}
        err, _warn = sf.check_sheet("author", sheets["author"], assigned)
        text = "\n".join(err)
        assert "根拠(頁・引用)が空" in text and "除外理由が語彙外" in text, text
        r2 = sheets["author"][ws.cell(2, 1).value]
        assert r2["decision"] == "Exclude" and r2["reason"] == REASON["I_scale"], r2   # 自動の判定・理由
        r3 = sheets["author"][ws.cell(3, 1).value]
        assert r3["decision"] == sf.expected_decision(r3)
        print(f"OK 検査: 注入した 2 種のミスをエラーとして検出({len(err)} 件)。判定・除外理由は自動で入る")

        # 注入ミスを直して、集計・協議・最終判定まで
        subprocess.run(PY + [str(HERE / "make_fulltext_sheets.py"), "--out", str(base), "--force"], check=True, capture_output=True)
        fill(base)
        r = subprocess.run(PY + [str(HERE / "score_fulltext.py"), "--dir", str(base)], capture_output=True, text=True, encoding="utf-8")
        assert "平均 κ" in r.stdout and r.returncode == 0, r.stdout[-800:] + r.stderr[-400:]
        wl = base / "consensus_worklist.csv"
        rows = list(csv.DictReader(wl.open(encoding="utf-8-sig", newline="")))
        for i, row in enumerate(rows):
            row["final_decision"] = "Include" if i % 2 == 0 else "Exclude"
            row["final_reason"] = "その他" if i % 2 else ""
        with wl.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        r = subprocess.run(PY + [str(HERE / "score_fulltext.py"), "--dir", str(base), "--final"], capture_output=True, text=True, encoding="utf-8")
        assert (base / "final_decisions.csv").exists() and (base / "extraction_table.csv").exists(), r.stdout[-800:]
        print(f"OK 集計: κ・協議リスト({len(rows)} 件)・最終判定・抽出表を書き出せた")
    print("全部 OK")


if __name__ == "__main__":
    main()
