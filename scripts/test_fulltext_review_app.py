#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_fulltext_review_app.py — 全文評価の閲覧アプリと反映スクリプトを、一時フォルダで端から端まで検証する
================================================================================
本物の判定シートには触れない。次を確認する:
  1. アプリ(HTML)が生成でき、Chrome/Edge のヘッドレスで開いて、判定の規則の自己テスト(#selftest)が通る
  2. 進捗 JSON を作り、apply_fulltext_review.py でシートへ反映できる
     (判定・自動の除外理由は式のまま残り、選び直した理由・抽出列だけが値として入る)
  3. 反映後のシートを score_fulltext.check_sheet に通してエラーが出ない
  4. 他人の JSON・語彙外の値・既存セルの上書きは止まる
実行:  python -X utf8 scripts/test_fulltext_review_app.py
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from make_fulltext_sheets import COL_INDEX  # noqa: E402
import score_fulltext as sf  # noqa: E402

PY = [sys.executable, "-X", "utf8"]
BROWSERS = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"]


def run(args, ok=True):
    r = subprocess.run(PY + args, capture_output=True, text=True, encoding="utf-8")
    assert (r.returncode == 0) == ok, (args, r.stdout[-600:], r.stderr[-400:])
    return r.stdout + r.stderr


def main() -> None:
    with tempfile.TemporaryDirectory() as td:
        base = Path(td) / "phase4"
        run([str(HERE / "make_fulltext_sheets.py"), "--out", str(base)])
        out = run([str(HERE / "make_fulltext_review_app.py"), "--dir", str(base), "--id", "author"])
        html = base / "author" / "review_author.html"
        assert html.exists() and "100 件" in out, out
        print("OK 生成: review_author.html")

        browser = next((b for b in BROWSERS if Path(b).exists()), None)
        if browser:
            r = subprocess.run([browser, "--headless", "--disable-gpu", "--allow-file-access-from-files",
                                "--virtual-time-budget=4000", "--dump-dom", html.resolve().as_uri() + "#selftest"],
                               capture_output=True, text=True, encoding="utf-8", timeout=120)
            assert 'id="selftest">SELFTEST OK' in r.stdout, r.stdout[-500:]
            print("OK 自己テスト: ブラウザで判定規則(Include/Exclude/Unsure・自動の理由・要修正)が通った")
        else:
            print("SKIP 自己テスト: Chrome/Edge が見つからない")

        # --- 進捗 JSON を組み立てる(アプリが書き出す形) ---
        sheet = base / "author" / "phase4_sheet_author.xlsx"
        ws = openpyxl.load_workbook(sheet)["判定"]
        rows = [(ws.cell(r, 1).value, ws.cell(r, COL_INDEX["extractor"]).value) for r in range(2, 8)]
        ext_id = next(rid for rid, x in rows if x == "Y")
        other_ids = [rid for rid, _x in rows if rid != ext_id]
        ycrit = {k: "Y" for k in ("P", "I_hmd", "I_scale", "C", "O", "S")}
        entries = {
            ext_id: {"c": ycrit, "reason": "", "evidence": "", "memo": "",
                     "ext": {"n_total": "24", "population": "大学生", "device": "HTC Vive",
                             "scale_target": ["身体全体", "IPD"], "scale_factor": "0.3x", "non_visual": ["なし"],
                             "task": ["距離推定"], "outcome_measures": "距離推定(行動)", "main_finding": "結果",
                             "attribution_quote": "記載なし", "scale_direction": "縮小", "action_contingent": "N"}},
            other_ids[0]: {"c": {**ycrit, "O": "N"}, "reason": "", "evidence": 'p.4 "no measure"', "memo": "", "ext": {}},
            other_ids[1]: {"c": {**ycrit, "S": "N"}, "reason": "S: 原著論文でない", "evidence": "p.1 review",
                           "memo": "", "ext": {}},
            other_ids[2]: {"c": {**ycrit, "C": "N"}, "reason": "", "evidence": "p.2 single", "memo": "", "ext": {}},
        }
        prog = {"format": "phase4-progress", "version": 1, "sheet": "author", "fingerprint": "x",
                "saved_at": "2026-10-21T00:00:00Z", "counts": {}, "entries": entries}
        pj = base / "author" / "progress_author.json"
        pj.write_text(json.dumps(prog, ensure_ascii=False), encoding="utf-8")

        run([str(HERE / "apply_fulltext_review.py"), "--id", "author", "--input", str(pj), "--dir", str(base), "--dry-run"])
        out = run([str(HERE / "apply_fulltext_review.py"), "--id", "author", "--input", str(pj), "--dir", str(base)])
        assert "更新した" in out, out
        assert list((base / "author" / "backup").glob("*.xlsx")), "バックアップが無い"
        print("OK 反映: 検証 → 書き込み → 元のファイルを backup/ へ退避")

        ws = openpyxl.load_workbook(sheet)["判定"]
        rowof = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
        g = lambda rid, k: ws.cell(rowof[rid], COL_INDEX[k]).value  # noqa: E731
        assert str(g(ext_id, "decision")).startswith("="), "判定の式が消えた"
        assert str(g(other_ids[0], "reason")).startswith("="), "自動の理由の式が消えた"
        assert g(other_ids[1], "reason") == "S: 原著論文でない", "選び直した理由が入っていない"
        assert g(ext_id, "n_total") == 24 and g(ext_id, "scale_target") == "身体全体; IPD"
        assert g(other_ids[0], "evidence") == 'p.4 "no measure"'
        print("OK 内容: 判定・自動の理由は式のまま、選び直した理由・抽出列は値として入った")

        # 反映後のシートを採点側の検査に通す(式は Excel で保存しないと値が無いので、採点側は 6 列から計算する)
        rows_loaded = sf.load_sheet(sheet)
        assigned = {a["record_id"]: a for a in sf.load_assignment(base)}
        sub = {rid: rows_loaded[rid] for rid in entries}
        err, _w = sf.check_sheet("author", sub, assigned)
        assert not err, err
        assert sub[other_ids[0]]["reason"] == "O: スケール知覚の測定が無い"
        assert sub[other_ids[2]]["reason"] == "その他"
        print("OK 採点側: 反映後のシートに検査エラー無し。自動の理由(O / その他)が計算された")

        # --- 止まるべきもの ---
        out = run([str(HERE / "apply_fulltext_review.py"), "--id", "kataoka", "--input", str(pj), "--dir", str(base)], ok=False)
        assert "シート ID が違う" in out, out
        bad = {**prog, "entries": {ext_id: {"c": {"P": "X"}, "ext": {}}}}
        pj2 = base / "author" / "bad.json"
        pj2.write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
        out = run([str(HERE / "apply_fulltext_review.py"), "--id", "author", "--input", str(pj2), "--dir", str(base)], ok=False)
        assert "Y/N/? でない" in out, out
        changed = {**prog, "entries": {ext_id: {**entries[ext_id], "c": {**ycrit, "P": "N"}}}}
        pj3 = base / "author" / "conflict.json"
        pj3.write_text(json.dumps(changed, ensure_ascii=False), encoding="utf-8")
        out = run([str(HERE / "apply_fulltext_review.py"), "--id", "author", "--input", str(pj3), "--dir", str(base)], ok=False)
        assert "既に別の値" in out, out
        run([str(HERE / "apply_fulltext_review.py"), "--id", "author", "--input", str(pj3), "--dir", str(base), "--force"])
        print("OK 防御: 他人の JSON・語彙外・既存セルの上書き(--force なし)は止まり、--force で通る")
        # --- スマホ向け配信(serve_review_app.py): 合言葉・配るファイルの限定・自動保存 ---
        import time
        import urllib.error
        import urllib.request
        (base / "author" / "pdf").mkdir(exist_ok=True)
        (base / "author" / "pdf" / "t.pdf").write_bytes(b"%PDF-1.4 test")
        (base / "author" / "secret.txt").write_text("x", encoding="utf-8")
        port = 18765
        proc = subprocess.Popen(PY + [str(HERE / "serve_review_app.py"), "--id", "author", "--dir", str(base),
                                      "--port", str(port), "--token", "tok"],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            u = f"http://127.0.0.1:{port}/tok/"

            def req(path, data=None):
                r = urllib.request.Request(u + path, data=data, method="POST" if data is not None else "GET",
                                           headers={"Content-Type": "application/json"})
                try:
                    with urllib.request.urlopen(r, timeout=5) as resp:
                        return resp.status, resp.read()
                except urllib.error.HTTPError as e:
                    return e.code, b""
            for _ in range(30):
                try:
                    st, body = req("review_author.html")
                    break
                except OSError:
                    time.sleep(0.2)
            assert st == 200 and b"mnav" in body, "アプリが配信されない"
            assert req("pdf/t.pdf")[0] == 200
            assert req("secret.txt")[0] == 404 and req("../secret.txt")[0] == 404 and req("pdf/../secret.txt")[0] == 404
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{port}/review_author.html", timeout=5)
                raise AssertionError("合言葉なしで開けた")
            except urllib.error.HTTPError as e:
                assert e.code == 404
            assert req("save", json.dumps({**prog, "sheet": "kataoka"}).encode())[0] == 400, "他人のシート ID を受けた"
            assert req("save", b"not json")[0] == 400
            assert req("save", json.dumps(prog, ensure_ascii=False).encode())[0] == 200
            st, body = req("progress.json")
            assert st == 200 and json.loads(body)["sheet"] == "author"
        finally:
            proc.terminate()
            proc.wait(timeout=10)
        print("OK 配信: 合言葉なし・別ファイルは 404、他人のシート ID の保存は拒否、保存と取得が通る")

        if browser:                                    # スマホ幅でも自己テストが通り、タブ切り替えが動く
            r = subprocess.run([browser, "--headless", "--disable-gpu", "--allow-file-access-from-files",
                                "--window-size=390,800", "--virtual-time-budget=4000", "--dump-dom",
                                html.resolve().as_uri() + "#selftest"],
                               capture_output=True, text=True, encoding="utf-8", timeout=120)
            assert 'id="selftest">SELFTEST OK' in r.stdout and "m-form" in r.stdout, r.stdout[-400:]
            print("OK スマホ幅(390px): 自己テストが通る")
    print("全部 OK")


if __name__ == "__main__":
    main()
