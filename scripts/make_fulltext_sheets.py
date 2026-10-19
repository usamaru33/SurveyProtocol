#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_fulltext_sheets.py — Phase 4(全文評価)の割当表と判定シート(Excel)を生成する
================================================================================

【何を】
  1. 担当の割当表 `assignment.csv`(均等化 案 B。Rev.29)
  2. 評価者ごとの判定シート `<who>/phase4_sheet_<who>.xlsx`
     (判定の6列・判定・除外理由・根拠・メモ + 事実の抽出列17列。すべて Excel のプルダウンで選べる)
  3. `--copy-pdfs` を付けると、評価者ごとのフォルダ `<who>/pdf/` に担当分の全文 PDF をコピーする

【割当(Rev.29: 均等化 案 B)】
  対象 = fulltext/manifest.csv の全件(未取得も含む。未取得は配布しないだけで割当は決める)。
  record_id を昇順に整列し、3 ペア [著者×Kataoka, 著者×Watanabe, Kataoka×Watanabe] に順に割り当てる
  (149 件なら 50 / 50 / 49)。乱数を使わないので、誰が実行しても同じ結果になる。
  `extractor`(事実の抽出列を書く 1 名)は、各ペア内で record_id を整列して交互に割り当てる。
  ※ manifest.csv の `pair` 列は仮割当(md5)なので使わない。

【シート】
  - 担当分のうち**全文 PDF を取得済み**の記録だけを載せる(未取得は取得後に追加配布する)。
  - 他の評価者の判定列は存在しない(独立性)。担当ペアの相手の名前だけが載る。
  - 記入列以外はロック。判定の 6 列は Y / N / ? のプルダウン。

【入力】 fulltext/manifest.csv / screening/final/final_decisions.csv(題名・DOI の補完)/ data/merged/ResearchVR4.csv と screening/**/sheet_*.csv(掲載先・年。照合できた分だけ)
【出力】 screening/phase4/assignment.csv と screening/phase4/<who>/phase4_sheet_<who>.xlsx

実行:
  python -X utf8 scripts/make_fulltext_sheets.py                    # 割当表とシートを生成
  python -X utf8 scripts/make_fulltext_sheets.py --copy-pdfs        # 担当分の PDF も評価者フォルダへ
  python -X utf8 scripts/make_fulltext_sheets.py --out <dir> --manifest <csv>   # 検証用(別の場所に出力)
  既存のシートは上書きしない(記入済みを壊さないため)。作り直すなら --force。
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
ROOT = _HERE.parent
sys.path.insert(0, str(_HERE))

csv.field_size_limit(10 ** 9)

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("[ERROR] openpyxl が必要です:  pip install openpyxl")

from make_screening_sheets import REVIEWERS  # noqa: E402
from make_screening_xlsx import EXCLUDE_REASONS, REASON_VALUES, REASON_OTHER  # noqa: E402  (語彙の正は 1 箇所)

# --- 割当 -------------------------------------------------------------------
PAIRS = [("author", "kataoka"), ("author", "watanabe"), ("kataoka", "watanabe")]

# --- 判定 -------------------------------------------------------------------
YNQ = ["Y", "N", "?"]
DECISIONS = ["Include", "Exclude", "Unsure"]
# (キー, シートの見出し, 説明)
CRITERIA = [
    ("P", "P 対象者", "健常成人か。Y=健常成人 / N=小児・患者・臨床集団など / ?=記載なし・判断できない"),
    ("I_hmd", "I-1 HMD の VR", "HMD を使った没入型 VR か。デスクトップ・実環境・CAVE=N。AR/MR(see-through)は MTG で決定のため ? にしてメモに「MR」"),
    ("I_scale", "I-2 スケール関係", "自己または環境のスケールに関係する操作・刺激があるか。身体所有感の誘導・パーソナルスペース・触覚だけでスケールと無関係=N"),
    ("C", "C 比較対象", "比較条件があるか(操作の有無・倍率間・視覚単独と多感覚)。単一条件の記述研究=N"),
    ("O", "O 定量指標", "自己・環境のスケール変容を測る定量データがあるか(身体所有感の質問紙も Y)"),
    ("S", "S 実験", "人を対象にした客観的なユーザー実験か。総説・技術提案・シミュレーションのみ=N。学位論文も実験を報告していれば Y"),
]
CRIT_KEYS = [k for k, _, _ in CRITERIA]
CRIT_LABEL = {k: l for k, l, _ in CRITERIA}

# --- 事実の抽出列(Rev.29。形式の正は docs/reference/phase4_reviewer_guide.md §3.5) -------------
# kind: text(自由) / int(整数 or 記載なし) / vocab(1 つ選ぶ) / multi(語彙から複数。`;` 区切りで入力可) / ynq
SCALE_TARGET = ["身体全体", "身体部位", "視点の高さ", "IPD", "環境", "物体", "なし(操作せず測定のみ)"]
SCALE_DIRECTION = ["縮小", "拡大", "両方", "該当なし"]
NON_VISUAL = ["聴覚", "触覚", "運動同期(歩行・身体動作)", "前庭", "その他", "なし"]
TASK = ["距離推定", "サイズ推定・マッチング", "リーチ・把持", "歩行・移動",
        "アフォーダンス判断", "身体サイズ判断", "質問紙のみ", "その他"]
SOCIAL = ["単独", "他アバターあり", "記載なし"]
INDIV = ["報告のみ", "分析に使った", "記載なし"]
# (キー, 見出し, 必須, kind, 語彙, 幅, 説明)
EXTRACT = [
    ("n_total", "参加者数(全体)", True, "int", None, 12, "全参加者数(除外前)の整数。書かれていなければ「記載なし」"),
    ("n_analyzed", "分析に使った人数", False, "int", None, 12, "n_total と同じなら空でよい"),
    ("population", "属性・年齢・性別", True, "text", None, 30, "例: 大学生・大学院生; 22.1±2.3 歳; 女 12 / 男 12。無ければ「記載なし」"),
    ("device", "HMD の機種", True, "text", None, 24, "例: HTC Vive; Oculus Rift CV1。複数は ; で区切る。無ければ「記載なし」"),
    ("scale_target", "何のスケールを変えたか", True, "multi", SCALE_TARGET, 24,
     "選択肢から選ぶ。複数あるときは「身体全体; IPD」のように ; で区切って直接入力してよい"),
    ("scale_factor", "倍率・値", True, "text", None, 28, "例: 身体 0.3×・1×・3× / アイハイト 1.0・1.7・2.5 m。無ければ「記載なし」"),
    ("scale_direction", "縮小/拡大", False, "vocab", SCALE_DIRECTION, 12, "縮小 / 拡大 / 両方 / 該当なし"),
    ("non_visual", "視覚以外の感覚", True, "multi", NON_VISUAL, 24,
     "選択肢から選ぶ。複数あるときは ; で区切って直接入力してよい。無いときは「なし」"),
    ("non_visual_detail", "視覚以外の提示内容", False, "text", None, 28, "何をどう提示したか。例: 足音の高周波を増幅。「なし」なら空"),
    ("action_contingent", "動作に連動?", False, "ynq", YNQ, 11, "視覚以外の手がかりが参加者自身の動作に連動して生じるか"),
    ("task", "課題", True, "multi", TASK, 24, "選択肢から選ぶ。複数あるときは ; で区切って直接入力してよい"),
    ("outcome_measures", "従属変数", True, "text", None, 40,
     "「名称(種別)」を ; で並べる。種別=主観報告/質問紙/行動/生理。例: 物体サイズ推定(行動); 身体所有感 SoE(質問紙)"),
    ("social_context", "社会的文脈", False, "vocab", SOCIAL, 14, "単独 / 他アバターあり / 記載なし"),
    ("individual_diff", "個人差の扱い", False, "vocab", INDIV, 14, "身長・VR 経験・性別・年齢を: 報告のみ / 分析に使った / 記載なし"),
    ("null_effect", "効果なしの条件?", False, "ynq", YNQ, 11, "操作の効果が有意でなかった条件があるか"),
    ("main_finding", "主な結果(1文)", True, "text", None, 44, "著者が報告した主な結果と方向。日本語 1 文"),
    ("attribution_quote", "帰属の記述(原文引用)", True, "text", None, 48,
     "変化を自己の変容と環境の変容のどちらに帰属したかの記述を原文のまま。末尾に (p.N)。無ければ「記載なし」"),
]
EXTRACT_KEYS = [e[0] for e in EXTRACT]
EXTRACT_REQUIRED = [e[0] for e in EXTRACT if e[2]]

# --- 配色(make_screening_xlsx.py と揃える。色だけに意味を持たせない) ---------------------
C_HEADER_BG = "1F3864"
C_HEADER_FG = "FFFFFF"
C_INPUT_BG = "FFF7E6"      # 判定の記入列(淡い橙)
C_EXTRACT_BG = "E8F1FB"    # 抽出列(淡い青)
C_BORDER = "BFBFBF"
C_WARN_BG = "FFC7CE"

LIST_SHEET = "_選択肢"

# 判定シートの列: (キー, 見出し, 幅, 折り返し, 区分)  区分: fixed / input / formula / extract
FIXED = [
    ("record_id", "ID", 13),
    ("extractor", "抽出担当", 9),
    ("partner", "担当ペアの相手", 16),
    ("title", "タイトル", 52),
    ("year", "年", 7),
    ("venue", "掲載先", 30),
    ("doi", "DOI", 26),
    ("pdf", "全文 PDF", 20),
]
INPUT_TAIL = [
    ("decision", "判定 ★", 12),
    ("reason", "除外理由 ★", 34),
    ("evidence", "根拠(頁・引用)", 40),
    ("memo", "メモ", 30),
]


def column_plan() -> list[tuple[str, str, int, str]]:
    cols = [(k, h, w, "fixed") for k, h, w in FIXED]
    cols += [(k, f"{CRIT_LABEL[k]} ★", 14, "input") for k in CRIT_KEYS]
    cols += [("rule_decision", "規則の判定(自動)", 14, "formula")]
    cols += [(k, h, w, "input") for k, h, w in INPUT_TAIL]
    cols += [(k, f"{h}{' ●' if req else ''}", w, "extract") for k, h, req, _kd, _v, w, _d in EXTRACT]
    return cols


COLS = column_plan()
COL_INDEX = {k: i for i, (k, _h, _w, _t) in enumerate(COLS, start=1)}


def col_letter(key: str) -> str:
    return get_column_letter(COL_INDEX[key])


# --- 割当 -------------------------------------------------------------------
def assign(record_ids: list[str]) -> dict[str, dict]:
    """record_id 昇順に 3 ペアへ順に割り当て、ペア内で抽出担当を交互にする。"""
    ids = sorted(record_ids)
    out: dict[str, dict] = {}
    per_pair: dict[int, list[str]] = {0: [], 1: [], 2: []}
    for i, rid in enumerate(ids):
        per_pair[i % 3].append(rid)
    for p, members in per_pair.items():
        a, b = PAIRS[p]
        for j, rid in enumerate(members):
            out[rid] = {"pair": f"{a}×{b}", "reviewer_a": a, "reviewer_b": b,
                        "extractor": a if j % 2 == 0 else b}
    return out


# --- 入力 -------------------------------------------------------------------
def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def norm_title(s: str) -> str:
    return re.sub(r"[^0-9a-z]+", " ", (s or "").lower()).strip()


def load_meta(path: Path) -> tuple[dict, dict]:
    """ResearchVR4 から DOI / 正規化題名 → (掲載先, 年)。無ければ空の辞書。"""
    by_doi: dict[str, tuple[str, str]] = {}
    by_title: dict[str, tuple[str, str]] = {}
    if not path.exists():
        return by_doi, by_title
    for r in read_csv(path):
        v = (r.get("Publication Title", ""), r.get("Publication Year", ""))
        d = (r.get("DOI") or "").lower().strip()
        if d:
            by_doi.setdefault(d, v)
        by_title.setdefault(norm_title(r.get("Title", "")), v)
    return by_doi, by_title


def load_screening_meta(root: Path) -> dict[str, tuple[str, str]]:
    """Phase 3b の判定シート(screening/**/*sheet_*.csv)から record_id → (掲載先, 年)。
    ResearchVR4 に無い文献追跡由来の記録を補う。ファイルが無ければ空。"""
    out: dict[str, tuple[str, str]] = {}
    for f in sorted(root.glob("**/*sheet_*.csv")):
        if "phase4" in f.parts:
            continue
        try:
            for r in read_csv(f):
                rid, v = r.get("record_id", ""), r.get("venue", "")
                if rid and v:
                    out.setdefault(rid, (v, r.get("year", "")))
        except (OSError, csv.Error):
            continue
    return out


# --- Excel ------------------------------------------------------------------
def build_list_sheet(wb) -> dict[str, str]:
    """プルダウンの選択肢を置く隠しシート。名前 → DataValidation 用の範囲参照 を返す。"""
    ws = wb.create_sheet(LIST_SHEET)
    lists = {
        "ynq": YNQ, "decision": DECISIONS, "reason": REASON_VALUES,
        "scale_target": SCALE_TARGET, "scale_direction": SCALE_DIRECTION,
        "non_visual": NON_VISUAL, "task": TASK, "social_context": SOCIAL,
        "individual_diff": INDIV,
    }
    refs = {}
    for j, (name, vals) in enumerate(lists.items(), start=1):
        ws.cell(row=1, column=j, value=name)
        for i, v in enumerate(vals, start=2):
            ws.cell(row=i, column=j, value=v)
        L = get_column_letter(j)
        refs[name] = f"={LIST_SHEET}!${L}$2:${L}${len(vals) + 1}"
        ws.column_dimensions[L].width = 28
    ws.protection.sheet = True
    ws.sheet_state = "hidden"
    return refs


INTRO = [
    ("Phase 4  全文評価  判定シート", "title"),
    ("", ""),
    ("担当: {reviewer}", "lead"),
    ("担当件数: {n} 件(うち抽出担当 {ne} 件)", "lead"),
    ("", ""),
    ("■ やること(1 件ごと)", "h"),
    ("1. 「全文 PDF」列のリンクを開き、全文を読む(抄録だけで判定しない)", ""),
    ("2. 「判定」シートの ★ の 6 列(P・I-1・I-2・C・O・S)を Y / N / ? から選ぶ。すべてプルダウンです", ""),
    ("3. 「判定」を選ぶ: 6 列がすべて Y → Include / 1 つでも N → Exclude / N が無く ? がある → Unsure", ""),
    ("   「規則の判定(自動)」列に、規則どおりの判定が出ます。食い違うと「判定」が赤くなります", ""),
    ("4. Exclude のときは「除外理由」(プルダウン)と「根拠(頁・引用)」を必ず書く", ""),
    ("5. 迷った点は「メモ」に書く。基準の解釈に迷ったら、その基準を ? にして協議に回す", ""),
    ("6. 「抽出担当」が Y の論文で、自分の判定が Include か Unsure のときは、青い列(事実の抽出)を書く", ""),
    ("   ●が付いた列は必須です。論文に書かれていない事実は「記載なし」と書く(空欄にしない)", ""),
    ("", ""),
    ("■ 判定の考え方", "h"),
    ("Phase 3b と違い、迷いを Include に倒さなくて構いません。全文評価は最終判断で、", ""),
    ("Include と Exclude のミスは同じ重さです。迷いは ? / Unsure で正直に出してください。", ""),
    ("Unsure は協議で必ず Include か Exclude に決めます(最終判定には残りません)。", ""),
    ("", ""),
    ("■ 各基準(要約。詳細は phase4_reviewer_guide.md)", "h"),
    ("__CRITERIA__", ""),
    ("", ""),
    ("■ 除外理由(1 件につき 1 つ。複数当てはまるときは上のものを選ぶ)", "h"),
    ("__REASONS__", ""),
    ("C(比較対象)が N のときは「その他」を選び、メモに「C: 比較対象が無い」と書く。", "lead"),
    ("", ""),
    ("■ 事実の抽出列", "h"),
    ("論文に書かれた事実だけを書く。Visual-Global や Self-scale などの分類の語彙は使わない。", ""),
    ("複数あるときは ; で区切る。引用は原文のまま、末尾に (p.N)。", ""),
    ("選択式の列(何のスケールを変えたか・視覚以外の感覚・課題)は、複数のときだけ手入力してかまいません。", ""),
    ("__EXTRACT__", ""),
    ("", ""),
    ("■ お願い", "h"),
    ("他の評価者のファイルは開かない・判定の内容を話さないでください(独立性。κ を算出します)。", "warn"),
    ("判定に LLM は使わないでください(読むための翻訳は構いません)。", "warn"),
    ("PDF が別の論文だった・開けない場合は、判定せず record_id と一緒に著者へ連絡してください。", ""),
    ("", ""),
    ("■ 使い方のヒント", "h"),
    ("・見出し行と ID 列は固定してあります。オートフィルタで「判定」を (空白) にすると未記入だけ表示できます。", ""),
    ("・記入列以外はロックしてあります。行の高さと列の幅は変えられます。", ""),
    ("・「全文 PDF」のリンクは、このファイルと同じフォルダの pdf/ フォルダを指します。", ""),
    ("", ""),
    ("■ 終わったら", "h"),
    ("このファイルを著者に渡してください。集計と κ の算出はスクリプト(score_fulltext.py)が行います。", ""),
]


def build_intro(ws, reviewer: str, n: int, ne: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 130
    styles = {
        "title": Font(size=16, bold=True, color=C_HEADER_BG),
        "h": Font(size=12, bold=True, color=C_HEADER_BG),
        "lead": Font(size=11, bold=True),
        "warn": Font(size=11, bold=True, color="9C0006"),
        "": Font(size=11),
    }
    lines: list[tuple[str, str]] = []
    for text, kind in INTRO:
        if text == "__CRITERIA__":
            lines += [(f"  ・{CRIT_LABEL[k]} … {d}", "") for k, _l, d in CRITERIA]
        elif text == "__REASONS__":
            lines += [(f"  ・{v} … {d}", "") for v, d in EXCLUDE_REASONS]
        elif text == "__EXTRACT__":
            lines += [(f"  ・{h}{' ●' if req else ''} … {d}", "") for _k, h, req, _kd, _v, _w, d in EXTRACT]
        else:
            lines.append((text.format(reviewer=reviewer, n=n, ne=ne) if text else "", kind))
    for i, (text, kind) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = styles[kind]
        c.alignment = Alignment(vertical="center", wrap_text=False)
        ws.row_dimensions[i].height = 24 if kind in ("title", "h") else 18
    ws.protection.sheet = True


def build_sheet(ws, rows: list[dict], refs: dict[str, str]) -> None:
    thin = Side(style="thin", color=C_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    n = len(rows)
    last = n + 1

    # --- 見出し ---
    for j, (key, label, width, typ) in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=j, value=label)
        c.font = Font(bold=True, color=C_HEADER_FG, size=11)
        c.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(j)].width = width
        helps = {k: d for k, _l, d in CRITERIA}
        helps.update({e[0]: e[6] for e in EXTRACT})
        if key in helps:
            c.comment = Comment(helps[key], "phase4 protocol")
            c.comment.width = 320
            c.comment.height = 110
    ws.row_dimensions[1].height = 44

    # --- 本体 ---
    in_fill = PatternFill("solid", fgColor=C_INPUT_BG)
    ex_fill = PatternFill("solid", fgColor=C_EXTRACT_BG)
    wrap_keys = {"title", "venue", "evidence", "memo", "population", "scale_factor", "non_visual_detail",
                 "outcome_measures", "main_finding", "attribution_quote", "reason", "device"}
    center_keys = {"year", "extractor", *CRIT_KEYS, "rule_decision", "decision",
                   "action_contingent", "null_effect", "scale_direction", "n_total", "n_analyzed"}
    crit_first, crit_last = col_letter(CRIT_KEYS[0]), col_letter(CRIT_KEYS[-1])
    for i, row in enumerate(rows, start=2):
        for j, (key, _l, _w, typ) in enumerate(COLS, start=1):
            if key == "rule_decision":
                rng = f"{crit_first}{i}:{crit_last}{i}"
                val = (f'=IF(COUNTIF({rng},"N")>0,"Exclude",IF(COUNTBLANK({rng})>0,"",'
                       f'IF(COUNTIF({rng},"~?")>0,"Unsure","Include")))')
            elif key == "year":
                try:
                    val = int(row.get("year", ""))
                except (TypeError, ValueError):
                    val = row.get("year", "")
            else:
                val = row.get(key, "") if typ == "fixed" else None
            c = ws.cell(row=i, column=j, value=val)
            c.border = border
            c.font = Font(size=11)
            c.alignment = Alignment(vertical="top", wrap_text=key in wrap_keys,
                                    horizontal="center" if key in center_keys else "left")
            if typ == "input":
                c.fill = in_fill
                c.protection = Protection(locked=False)
            elif typ == "extract":
                c.fill = ex_fill
                c.protection = Protection(locked=False)
            else:
                c.protection = Protection(locked=True)
        if row.get("pdf"):
            pc = ws.cell(row=i, column=COL_INDEX["pdf"])
            pc.hyperlink = f"pdf/{row['pdf']}"
            pc.font = Font(size=11, color="0563C1", underline="single")
        ws.row_dimensions[i].height = 60

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"

    # --- プルダウン ---
    def add_list(key: str, ref: str, *, title: str, prompt: str, strict: bool = True) -> None:
        L = col_letter(key)
        dv = DataValidation(type="list", formula1=ref, allow_blank=True, showDropDown=False)
        dv.errorStyle = "stop" if strict else "warning"
        dv.errorTitle = "入力できない値です" if strict else "一覧にない値です"
        dv.error = ("一覧から選んでください。" if strict else
                    "一覧から選ぶか、複数のときは「A; B」のように ; で区切って入力してください。")
        dv.promptTitle, dv.prompt = title, prompt
        ws.add_data_validation(dv)
        dv.add(f"{L}2:{L}{last}")

    for k in CRIT_KEYS:
        add_list(k, refs["ynq"], title=CRIT_LABEL[k], prompt="Y / N / ?(迷い・記載なし)")
    add_list("decision", refs["decision"], title="判定",
             prompt="6 列すべて Y=Include / 1 つでも N=Exclude / N が無く ? がある=Unsure")
    add_list("reason", refs["reason"], title="除外理由",
             prompt="Exclude のときは必須。複数当てはまるときは優先順位の上のもの1つ")
    for key, kind, vkey in (("scale_target", "multi", "scale_target"), ("non_visual", "multi", "non_visual"),
                            ("task", "multi", "task"), ("scale_direction", "vocab", "scale_direction"),
                            ("social_context", "vocab", "social_context"),
                            ("individual_diff", "vocab", "individual_diff")):
        add_list(key, refs[vkey], title=dict((e[0], e[1]) for e in EXTRACT)[key],
                 prompt="選択肢から選ぶ" + ("(複数のときは ; で区切って入力可)" if kind == "multi" else ""),
                 strict=(kind != "multi"))
    for key in ("action_contingent", "null_effect"):
        add_list(key, refs["ynq"], title=dict((e[0], e[1]) for e in EXTRACT)[key], prompt="Y / N / ?")

    # ※ COUNTIF の条件 "?" はワイルドカード(任意の 1 文字)なので、文字としての ? は "~?" と書く
    # --- 条件付き書式(提出前の自己チェック。色は補助で、文字でも分かる) ---------------------
    D, R, E, M = col_letter("decision"), col_letter("reason"), col_letter("evidence"), col_letter("memo")
    RD, X = col_letter("rule_decision"), col_letter("extractor")
    red = dict(fill=PatternFill("solid", fgColor=C_WARN_BG), font=Font(color="9C0006", bold=True))
    # (1) 判定が規則と食い違う
    ws.conditional_formatting.add(f"{D}2:{D}{last}", FormulaRule(
        formula=[f'AND({D}2<>"",{RD}2<>"",{D}2<>{RD}2)'], stopIfTrue=False, **red))
    # (2) Exclude なのに理由 / 根拠が空
    ws.conditional_formatting.add(f"{R}2:{R}{last}", FormulaRule(
        formula=[f'AND({D}2="Exclude",LEN(TRIM({R}2))=0)'], stopIfTrue=False, **red))
    ws.conditional_formatting.add(f"{E}2:{E}{last}", FormulaRule(
        formula=[f'AND({D}2="Exclude",LEN(TRIM({E}2))=0)'], stopIfTrue=False, **red))
    # (3) 理由が「その他」なのにメモが空
    ws.conditional_formatting.add(f"{M}2:{M}{last}", FormulaRule(
        formula=[f'AND({R}2="{REASON_OTHER}",LEN(TRIM({M}2))=0)'], stopIfTrue=False, **red))
    # (4) 抽出担当で、自分の判定が Include / Unsure なのに必須の抽出列が空
    for key in EXTRACT_REQUIRED:
        L = col_letter(key)
        ws.conditional_formatting.add(f"{L}2:{L}{last}", FormulaRule(
            formula=[f'AND({X}2="Y",OR({D}2="Include",{D}2="Unsure"),LEN(TRIM({L}2))=0)'],
            stopIfTrue=False, **red))
    # (5) 抽出担当でない行の抽出列を灰色にする(書かなくてよい)
    first_ex, last_ex = col_letter(EXTRACT_KEYS[0]), col_letter(EXTRACT_KEYS[-1])
    ws.conditional_formatting.add(f"{first_ex}2:{last_ex}{last}", FormulaRule(
        formula=[f'{X}2<>"Y"'], fill=PatternFill("solid", fgColor="E7E6E6"), stopIfTrue=False))

    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def main() -> None:
    ap = argparse.ArgumentParser(description="Phase 4 の割当表と判定シート(Excel)を生成")
    ap.add_argument("--manifest", type=Path, default=ROOT / "fulltext" / "manifest.csv")
    ap.add_argument("--final", type=Path, default=ROOT / "screening" / "final" / "final_decisions.csv")
    ap.add_argument("--meta", type=Path, default=ROOT / "data" / "merged" / "ResearchVR4.csv")
    ap.add_argument("--pdf-dir", type=Path, default=ROOT / "fulltext" / "pdf")
    ap.add_argument("--out", type=Path, default=ROOT / "screening" / "phase4")
    ap.add_argument("--copy-pdfs", action="store_true", help="担当分の PDF を評価者フォルダの pdf/ にコピー")
    ap.add_argument("--only", type=str, default="", help="特定の評価者だけ(author / kataoka / watanabe)")
    ap.add_argument("--force", action="store_true", help="既存のシートを上書き(記入済みの内容は失われる)")
    args = ap.parse_args()

    manifest = read_csv(args.manifest)
    ids = [r["record_id"] for r in manifest]
    if len(set(ids)) != len(ids):
        sys.exit("[ERROR] manifest.csv に record_id の重複がある")
    plan = assign(ids)

    final = {r["record_id"]: r for r in read_csv(args.final)} if args.final.exists() else {}
    by_doi, by_title = load_meta(args.meta)
    by_rid = load_screening_meta(ROOT / "screening")

    info: dict[str, dict] = {}
    for m in manifest:
        rid = m["record_id"]
        title = m["title"] or final.get(rid, {}).get("title", "")
        doi = m["doi"] or final.get(rid, {}).get("doi", "")
        venue, year = (by_doi.get((doi or "").lower()) or by_title.get(norm_title(title))
                       or by_rid.get(rid) or ("", ""))
        has_pdf = m["status"] == "ok" and (args.pdf_dir / f"{rid}.pdf").exists()
        info[rid] = {"record_id": rid, "title": title, "doi": doi, "venue": venue, "year": year,
                     "pdf": f"{rid}.pdf" if has_pdf else "", "status": m["status"], "has_pdf": has_pdf}

    # --- 割当表 ---
    args.out.mkdir(parents=True, exist_ok=True)
    apath = args.out / "assignment.csv"
    with apath.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["record_id", "pair", "reviewer_a", "reviewer_b", "extractor", "distributed", "doi", "title"])
        for rid in sorted(ids):
            p = plan[rid]
            w.writerow([rid, p["pair"], p["reviewer_a"], p["reviewer_b"], p["extractor"],
                        "Y" if info[rid]["has_pdf"] else "N", info[rid]["doi"], info[rid]["title"]])
    from collections import Counter
    pc = Counter(p["pair"] for p in plan.values())
    ev = Counter()
    for p in plan.values():
        ev[p["reviewer_a"]] += 1
        ev[p["reviewer_b"]] += 1
    print(f"[INFO] 割当表: {apath}  {len(ids)} 件  ペア別 {dict(pc)}  評価回数 {dict(ev)}")

    # --- シート ---
    targets = [args.only] if args.only else list(REVIEWERS)
    for who in targets:
        if who not in REVIEWERS:
            sys.exit(f"[ERROR] 未知の評価者: {who}(有効: {', '.join(REVIEWERS)})")
        mine = [rid for rid in sorted(ids) if who in (plan[rid]["reviewer_a"], plan[rid]["reviewer_b"])]
        dist = [rid for rid in mine if info[rid]["has_pdf"]]
        rows = []
        for rid in dist:
            p = plan[rid]
            partner = p["reviewer_b"] if p["reviewer_a"] == who else p["reviewer_a"]
            rows.append({**info[rid], "extractor": "Y" if p["extractor"] == who else "N",
                         "partner": REVIEWERS[partner]})
        d = args.out / who
        d.mkdir(parents=True, exist_ok=True)
        dst = d / f"phase4_sheet_{who}.xlsx"
        if dst.exists() and not args.force:
            print(f"[SKIP] {dst} は既にある(記入済みを壊さないため生成しない。作り直すなら --force)")
        else:
            wb = Workbook()
            intro = wb.active
            intro.title = "はじめに"
            build_intro(intro, REVIEWERS[who], len(rows), sum(1 for r in rows if r["extractor"] == "Y"))
            ws = wb.create_sheet("判定")
            refs = build_list_sheet(wb)
            build_sheet(ws, rows, refs)
            wb.active = 1
            wb.save(dst)
            print(f"[INFO] {dst}  {len(rows)} 行(割当 {len(mine)} 件のうち未取得 {len(mine) - len(dist)} 件は載せない)"
                  f"  抽出担当 {sum(1 for r in rows if r['extractor'] == 'Y')} 件")
        if args.copy_pdfs:
            pdir = d / "pdf"
            pdir.mkdir(exist_ok=True)
            for rid in dist:
                shutil.copyfile(args.pdf_dir / f"{rid}.pdf", pdir / f"{rid}.pdf")
            print(f"[INFO] PDF {len(dist)} 件を {pdir} にコピー")

    print("\n[NEXT] 記入中のシート(phase4_sheet_*)は .gitignore 対象。各評価者に自分のフォルダだけを渡す。"
          "\n       返却後は score_fulltext.py で検査・集計する。")


if __name__ == "__main__":
    main()
