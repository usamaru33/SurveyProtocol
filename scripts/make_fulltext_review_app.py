#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Phase 4(全文評価)用のローカル閲覧アプリを生成する。

Phase 3b の `make_review_app.py` の全文版。**判定は人が下す。** このアプリは「読みながら、
その場で記入する」ための道具で、判定(Include/Exclude/Unsure)と除外理由は
6 つの基準(Y/N/?)から**自動で計算して表示する**だけ(シートの式と同じ規則)。

画面:  左 = 全文 PDF / 右 = 記入欄(6 基準・判定(自動)・除外理由・根拠・メモ・事実の抽出)
入力:  screening/phase4/<id>/phase4_sheet_<id>.xlsx   (記入済みがあれば取り込む)
        screening/phase4/<id>/pdf/<論文名>.pdf         (HTML と同じフォルダに置く。相対リンクで開く)
出力:  screening/phase4/<id>/review_<id>.html           単体で開ける(依存なし・オフライン)

使い方:
    python -X utf8 scripts/make_fulltext_review_app.py                 # 3 人分すべて
    python -X utf8 scripts/make_fulltext_review_app.py --id author     # 1 人分
    python -X utf8 scripts/make_fulltext_review_app.py --id author --artifact   # スマホ単体用

`--artifact` は PDF を同梱しない 1 ファイル版(`review_<id>_artifact.html`)。PC もサーバーも要らず、
Claude のアーティファクトやスマホのブラウザでそのまま開ける。論文は オープンアクセスの URL
(`fulltext/manifest.csv` の source)か DOI のリンクで別タブに開く。進捗はブラウザ内に保存し、
「進捗を保存」で JSON を表示・コピー、「進捗を読み込む」で貼り付けて戻す(ファイル操作が要らない)。

HTML をブラウザ(Chrome / Edge)で開いて記入する。入力は逐次ブラウザに保存され、`S` で進捗 JSON を
書き出せる(「自動保存先を選ぶ」を 1 回押すと、以後は変更のたびにそのファイルへ書き込む)。
JSON は `scripts/apply_fulltext_review.py` で判定シート(xlsx)へ反映する。
**xlsx に直接書いても構わない**(このアプリは必須ではない)。

独立性: HTML には**自分の担当分だけ**が入る(担当ペアの相手の名前は出るが、相手の判定は存在しない)。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
ROOT = _HERE.parent
sys.path.insert(0, str(_HERE))

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("[ERROR] openpyxl が必要です:  pip install openpyxl")

from make_fulltext_sheets import (  # noqa: E402  (語彙・列・自動規則の正はシート生成側)
    AUTO_REASON_ORDER, COLS, COL_INDEX, CRITERIA, CRIT_KEYS, EXTRACT, REASON_OTHER, REASON_VALUES,
)
from make_screening_sheets import REVIEWERS  # noqa: E402

PHASE4 = ROOT / "screening" / "phase4"


def fingerprint(ids: list[str]) -> str:
    """進捗 JSON が同じ担当分に対するものかを照合する指紋(件数:ID 列の SHA-1 先頭 7 桁)。"""
    return "{}:{}".format(len(ids), hashlib.sha1("\n".join(ids).encode("utf-8")).hexdigest()[:7])


def load_records(who: str, base: Path) -> list[dict]:
    path = base / who / f"phase4_sheet_{who}.xlsx"
    if not path.exists():
        sys.exit(f"[ERROR] {path} が無い。先に make_fulltext_sheets.py を実行すること。")
    ws = load_workbook(path)["判定"]            # 式は式のまま読む(自動の列と、手で選んだ値を区別するため)
    header = [c.value for c in ws[1]]
    expect = [h for _k, h, _w, _t in COLS]
    if header != expect:
        sys.exit(f"[ERROR] {path.name} の列がこのスクリプトの想定と違う(シートを作り直すか、スクリプトを更新)。")
    kinds = {e[0]: (e[3], e[4]) for e in EXTRACT}
    recs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {k: ("" if v is None else v) for (k, *_), v in zip(COLS, row)}
        if not d["record_id"]:
            continue
        e = {"c": {k: str(d[k]).strip() for k in CRIT_KEYS if str(d[k]).strip()},
             "reason": "" if str(d["reason"]).startswith("=") else str(d["reason"]).strip(),
             "evidence": str(d["evidence"]).strip(), "memo": str(d["memo"]).strip(), "ext": {}}
        for k, (kind, vocab) in kinds.items():
            v = d[k]
            if v == "":
                continue
            if kind == "multi":
                e["ext"][k] = [t.strip() for t in str(v).replace("；", ";").split(";") if t.strip() in vocab]
            elif kind == "int":
                e["ext"][k] = str(int(v)) if isinstance(v, float) and v == int(v) else str(v).strip()
            else:
                e["ext"][k] = str(v).strip()
        touched = e["c"] or e["reason"] or e["evidence"] or e["memo"] or e["ext"]
        recs.append({"id": d["record_id"], "extractor": d["extractor"], "partner": d["partner"],
                     "title": d["title"], "year": d["year"], "venue": d["venue"], "doi": d["doi"],
                     "pdf": d["pdf"], "init": e if touched else None})
    return recs


def add_links(recs: list[dict]) -> None:
    """--artifact 用: PDF の代わりに開くリンク(OA の URL > DOI > Google Scholar のタイトル検索)。"""
    import csv
    from urllib.parse import quote
    src = {}
    with open(ROOT / "fulltext" / "manifest.csv", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            src[row["record_id"]] = row.get("source", "")
    for r in recs:
        oa = src.get(r["id"], "")
        r["url"] = (oa if oa.startswith("http") else
                    f"https://doi.org/{r['doi']}" if r["doi"] else
                    "https://scholar.google.com/scholar?q=" + quote(str(r["title"])))
        r["oa"] = oa.startswith("http")


def build_data(who: str, recs: list[dict], artifact: bool = False) -> dict:
    return {"artifact": artifact,
        "sheet": who, "reviewer": REVIEWERS[who], "fingerprint": fingerprint([r["id"] for r in recs]),
        "records": recs,
        "criteria": [{"k": k, "label": l, "desc": d} for k, l, d in CRITERIA],
        "extract": [{"k": k, "label": l, "req": req, "kind": kind, "vocab": vocab or [], "desc": desc}
                    for k, l, req, kind, vocab, _w, desc in EXTRACT],
        "autoReason": [[k, r] for k, r in AUTO_REASON_ORDER],
        "reasons": REASON_VALUES, "other": REASON_OTHER,
    }


def render(data: dict) -> str:
    blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    return TEMPLATE.replace("__DATA__", blob).replace("__WHO__", data["reviewer"])


def main() -> None:
    ap = argparse.ArgumentParser(description="Phase 4 全文評価の閲覧アプリ(HTML)を生成")
    ap.add_argument("--id", default="", help="author / kataoka / watanabe(省略で 3 人分)")
    ap.add_argument("--dir", type=Path, default=PHASE4)
    ap.add_argument("--artifact", action="store_true", help="PDF を同梱しないスマホ単体版を作る")
    args = ap.parse_args()
    targets = [args.id] if args.id else list(REVIEWERS)
    for who in targets:
        if who not in REVIEWERS:
            sys.exit(f"[ERROR] 未知の評価者: {who}(有効: {', '.join(REVIEWERS)})")
        recs = load_records(who, args.dir)
        if args.artifact:
            add_links(recs)
        out = args.dir / who / (f"review_{who}_artifact.html" if args.artifact else f"review_{who}.html")
        out.write_text(render(build_data(who, recs, args.artifact)), encoding="utf-8")
        done = sum(1 for r in recs if r["init"])
        print(f"[OK] {out}  {len(recs)} 件(うち記入あり {done} 件)")


TEMPLATE = r"""<!doctype html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Phase 4 全文評価 — __WHO__</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--fg:#1b1f24;--mut:#5b6572;--line:#d5d9df;--acc:#1f3864;--ok:#1a7f37;--bad:#b42318;--warn:#9a6700;
--y:#c6efce;--n:#f8cbad;--q:#ffeb9c}
*{box-sizing:border-box}
body{margin:0;font:14px/1.55 "Yu Gothic UI","Meiryo",sans-serif;background:var(--bg);color:var(--fg);height:100vh;display:flex;flex-direction:column}
header{display:flex;gap:12px;align-items:center;padding:6px 12px;background:var(--acc);color:#fff;flex-wrap:wrap}
header b{font-size:15px}
header .sp{flex:1}
header button{background:#fff;color:var(--acc);border:0;border-radius:4px;padding:3px 10px;cursor:pointer;font-size:13px}
#bar{height:6px;background:#dde3ee}#bar i{display:block;height:100%;background:var(--ok);width:0}
main{flex:1;display:flex;min-height:0}
#left{flex:1;min-width:0;display:flex;flex-direction:column;border-right:1px solid var(--line)}
#left iframe{flex:1;border:0;width:100%;background:#fff}
#nopdf{padding:24px;color:var(--mut)}
#right{width:520px;max-width:55vw;overflow-y:auto;padding:10px 14px;background:var(--card)}
.hide-pdf #left{display:none}.hide-pdf #right{width:100%;max-width:none}
h1{font-size:16px;margin:4px 0}
.meta{color:var(--mut);font-size:12px}
.nav{display:flex;gap:6px;align-items:center;margin:6px 0;flex-wrap:wrap}
.nav button,.nav select{padding:3px 8px;font-size:13px}
.sec{border:1px solid var(--line);border-radius:6px;padding:8px 10px;margin:10px 0;background:#fff}
.sec h2{font-size:13px;margin:0 0 6px;color:var(--acc)}
.crit{display:flex;align-items:center;gap:6px;padding:3px 4px;border-radius:4px}
.crit.on{background:#eef3fb;outline:2px solid #9db7e5}
.crit .lab{flex:1}.crit .lab small{color:var(--mut);display:block;font-size:11px;line-height:1.3}
.k{display:inline-block;min-width:18px;text-align:center;background:#e8ebf0;border-radius:3px;font-size:11px;margin-right:4px}
.seg button{border:1px solid var(--line);background:#fff;padding:3px 12px;cursor:pointer;font-size:13px}
.seg button:first-child{border-radius:4px 0 0 4px}.seg button:last-child{border-radius:0 4px 4px 0}
.seg button.on[data-v=Y]{background:var(--y);font-weight:700}.seg button.on[data-v=N]{background:var(--n);font-weight:700}.seg button.on[data-v="?"]{background:var(--q);font-weight:700}
#decision{font-size:18px;font-weight:700;padding:2px 12px;border-radius:4px;display:inline-block;min-width:110px;text-align:center;background:#eee}
#decision.Include{background:var(--y)}#decision.Exclude{background:var(--n)}#decision.Unsure{background:var(--q)}
label.f{display:block;font-size:12px;color:var(--mut);margin:6px 0 2px}
textarea,input[type=text],select{width:100%;font:inherit;padding:4px 6px;border:1px solid var(--line);border-radius:4px}
textarea{min-height:54px;resize:vertical}
.req{color:var(--bad)}
.chk label{display:inline-block;margin:2px 10px 2px 0;white-space:nowrap}
.quick{margin-left:4px;font-size:11px;padding:1px 6px}
#probs{color:var(--bad);font-size:12px;margin-top:4px}
#probs.okay{color:var(--ok)}
.dim{opacity:.55}
#legend{font-size:11px;color:var(--mut);margin-top:10px;line-height:1.7}
#list{position:fixed;inset:44px 0 0 0;background:rgba(255,255,255,.97);overflow:auto;padding:12px;display:none;z-index:5}
#list.on{display:block}
#list table{border-collapse:collapse;width:100%;font-size:12px}
#list td,#list th{border-bottom:1px solid var(--line);padding:3px 6px;text-align:left}
#list tr.row{cursor:pointer}#list tr.row:hover{background:#eef3fb}
#toast{position:fixed;right:14px;bottom:14px;background:#222;color:#fff;padding:6px 12px;border-radius:6px;display:none;z-index:9}
#selftest{display:none}
#jbox{display:none;position:fixed;inset:0;background:rgba(0,0,0,.45);z-index:20;align-items:center;justify-content:center;padding:12px}
#jbox>div{background:#fff;border-radius:8px;padding:12px;width:100%;max-width:640px}
#jbox textarea{height:40vh;font:12px monospace;margin:8px 0}#jbox button{padding:10px 18px;font-size:15px}
#mnav,#pdfopen{display:none}
@media (max-width:820px){
  body{height:100vh;height:100dvh;font-size:15px}
  header{padding:4px 8px;gap:4px 8px}header b{font-size:13px}
  header #stats{order:9;width:100%;font-size:11px}header #saved{font-size:11px}
  header button{padding:6px 10px;font-size:13px}#bPdf,#bLoad{display:none}
  main{flex-direction:column}
  #left,#right{width:100%;max-width:none;border:0;flex:1;min-height:0}
  #right{padding:8px 10px}
  body.m-form #left{display:none}body.m-pdf #right{display:none}
  #pdfopen{display:block;margin:12px;padding:16px;text-align:center;background:var(--acc);color:#fff;border-radius:8px;font-size:17px;text-decoration:none}
  #mnav{display:flex;gap:6px;padding:6px 8px calc(6px + env(safe-area-inset-bottom));background:#fff;border-top:1px solid var(--line)}
  #mnav button{flex:1;padding:12px 0;font-size:16px;border:1px solid var(--line);border-radius:8px;background:#f1f3f7}
  #mnav button.on{background:var(--acc);color:#fff;border-color:var(--acc)}
  #mnav #mPrev,#mnav #mNext{flex:0 0 52px}
  #right .nav #prev,#right .nav #next{display:none}
  .crit{flex-wrap:wrap}.crit .lab{flex:1 0 100%}.crit .seg{flex:1 0 100%;display:flex}
  .seg button,.seg button:first-child,.seg button:last-child{flex:1;padding:12px 0;font-size:17px;border-radius:0}
  .seg button:first-child{border-radius:8px 0 0 8px}.seg button:last-child{border-radius:0 8px 8px 0}
  textarea,input[type=text],select{font-size:16px;padding:8px}
  .nav select{padding:6px}
  .chk label{margin:6px 14px 6px 0;font-size:15px}.chk input{width:20px;height:20px;vertical-align:middle}
  .quick{padding:6px 10px;font-size:13px}
  #legend{display:none}
  #list{inset:0;top:84px}
  #toast{left:10px;right:10px;bottom:70px;text-align:center}
}
</style></head><body>
<header>
  <b>Phase 4 全文評価 — __WHO__</b>
  <span id="stats"></span><span class="sp"></span>
  <span id="saved" class="meta" style="color:#cfd8ea"></span>
  <button id="bSave" title="S キー">進捗を保存</button>
  <button id="bLoad" title="L キー">進捗を読み込む</button>
  <button id="bList" title="Tab キー">一覧</button>
  <button id="bPdf" title="G キー">PDF 表示/非表示</button>
</header>
<div id="bar"><i></i></div>
<main id="main">
  <section id="left"><a id="pdfopen" target="_blank" href="#">PDF を開く(別タブ)</a><iframe id="pdf" title="全文 PDF"></iframe><div id="nopdf" style="display:none"></div></section>
  <section id="right"></section>
</main>
<nav id="mnav"><button id="mPrev">◀</button><button id="mTabF" class="on">✏ 記入</button><button id="mTabP">📄 論文 PDF</button><button id="mNext">▶</button></nav>
<div id="list"></div>
<div id="toast"></div>
<input type="file" id="file" accept=".json,application/json" style="display:none">
<div id="selftest"></div>
<div id="jbox"><div><b id="jtitle"></b><textarea id="jtext"></textarea><p><button id="jgo"></button> <button id="jclose">閉じる</button></p></div></div>
<script type="application/json" id="data">__DATA__</script>
<script>
"use strict";
const D=JSON.parse(document.getElementById("data").textContent);
const CR=D.criteria, EX=D.extract;
const KEY="phase4:"+D.sheet+":"+D.fingerprint;
// アーティファクトの枠内などで localStorage が使えないときは、メモリに置く(閉じると消えるので JSON で保存)
const LS=(()=>{try{localStorage.setItem("_t","1");return localStorage}catch(e){const m={};return{getItem:k=>k in m?m[k]:null,setItem:(k,v)=>{m[k]=String(v)}}}})();
let store={};
try{store=JSON.parse(LS.getItem(KEY)||"{}")}catch(e){store={}}
for(const r of D.records){ if(!store[r.id]&&r.init) store[r.id]=JSON.parse(JSON.stringify(r.init)); }
let cur=0, view=[], fc=0, dirty=false, handle=null, filter="all", pdfOn=true, saveTimer=null, noSave=false;

const $=id=>document.getElementById(id);
const esc=s=>String(s==null?"":s).replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function E(id){return store[id]||(store[id]={c:{},reason:"",evidence:"",memo:"",ext:{}})}
function dec(e){const v=CR.map(x=>e.c[x.k]||"");if(v.includes("N"))return"Exclude";if(v.includes(""))return"";return v.includes("?")?"Unsure":"Include"}
function autoR(e){if(dec(e)!=="Exclude")return"";for(const [k,r] of D.autoReason){if(e.c[k]==="N")return r}return""}
function reasonOf(e){return dec(e)==="Exclude"?(e.reason||autoR(e)):""}
function filled(e,x){const v=e.ext[x.k];return Array.isArray(v)?v.length>0:(v!==undefined&&String(v).trim()!=="")}
function needExt(r,e){const d=dec(e);return r.extractor==="Y"&&(d==="Include"||d==="Unsure")}
function problems(r,e){
  const p=[],d=dec(e);
  if(d==="Exclude"){if(!(e.evidence||"").trim())p.push("根拠(頁・引用)")}
  if(reasonOf(e)===D.other&&e.c.C!=="N"&&!(e.memo||"").trim())p.push("メモ(理由が「その他」)");
  if(needExt(r,e))for(const x of EX)if(x.req&&!filled(e,x))p.push(x.label);
  for(const k of["n_total","n_analyzed"]){const v=e.ext[k];
    if(v!==undefined&&String(v).trim()!==""&&!(/^\d+$/.test(String(v).trim())||v==="記載なし"))p.push(k==="n_total"?"参加者数は整数か「記載なし」":"分析に使った人数は整数か「記載なし」")}
  return p}
function touched(e){return Object.keys(e.c).length||e.reason||e.evidence||e.memo||Object.keys(e.ext).length}

function rebuild(){
  const id=view[cur]!==undefined?D.records[view[cur]].id:null;
  view=[];
  D.records.forEach((r,i)=>{const e=store[r.id]||{c:{},ext:{}},d=dec(e),p=problems(r,E(r.id));
    const ok={all:1,todo:d==="",Exclude:d==="Exclude",Unsure:d==="Unsure",Include:d==="Include",
      ext:needExt(r,e)&&p.length>0,prob:p.length>0}[filter];
    if(ok)view.push(i)});
  const j=id?view.findIndex(i=>D.records[i].id===id):-1;
  cur=j>=0?j:Math.min(cur,Math.max(view.length-1,0));
}
function stats(){
  let done=0,inc=0,exc=0,uns=0,ext=0,extDone=0;
  for(const r of D.records){const e=store[r.id]||{c:{},ext:{}},d=dec(e);
    if(d){done++;if(d==="Include")inc++;if(d==="Exclude")exc++;if(d==="Unsure")uns++}
    if(needExt(r,e)){ext++;if(!problems(r,E(r.id)).length)extDone++}}
  return {done,inc,exc,uns,ext,extDone,n:D.records.length}}

function draw(){
  rebuild();
  const R=$("right");
  if(!view.length){R.innerHTML='<p>この絞り込みに該当する論文はありません。</p><div class="nav"><select id="flt">'+fltOpts()+'</select></div>';bindFlt();$("pdf").src="about:blank";return}
  const r=D.records[view[cur]],e=E(r.id);
  const crit=CR.map((x,i)=>'<div class="crit'+(i===fc?" on":"")+'" data-i="'+i+'"><span class="lab"><span class="k">'+(i+1)+'</span>'+esc(x.label)+' ★<small>'+esc(x.desc)+'</small></span>'
    +'<span class="seg" data-k="'+x.k+'">'+["Y","N","?"].map(v=>'<button data-v="'+v+'" class="'+(e.c[x.k]===v?"on":"")+'">'+v+'</button>').join("")+'</span></div>').join("");
  const exHtml=EX.map(x=>{
    const v=e.ext[x.k],lab='<label class="f">'+esc(x.label)+(x.req?' <span class="req">●</span>':'')+' <span style="color:#8a94a3">'+esc(x.desc)+'</span></label>';
    if(x.kind==="multi"){const set=new Set(Array.isArray(v)?v:[]);
      return lab+'<div class="chk" data-ms="'+x.k+'">'+x.vocab.map(o=>'<label><input type="checkbox" value="'+esc(o)+'"'+(set.has(o)?" checked":"")+'> '+esc(o)+'</label>').join("")+'</div>'}
    if(x.kind==="vocab")return lab+'<select data-ext="'+x.k+'"><option value="">(未選択)</option>'+x.vocab.map(o=>'<option'+(v===o?" selected":"")+'>'+esc(o)+'</option>').join("")+'</select>';
    if(x.kind==="ynq")return lab+'<span class="seg" data-ek="'+x.k+'">'+["Y","N","?"].map(o=>'<button data-v="'+o+'" class="'+(v===o?"on":"")+'">'+o+'</button>').join("")+'</span>';
    const big=["population","outcome_measures","main_finding","attribution_quote","scale_factor","non_visual_detail"].includes(x.k);
    const q=(x.kind==="int"||x.req)?'<button class="quick" data-fill="'+x.k+'" data-v="記載なし">記載なし</button>':"";
    return lab+(big?'<textarea data-ext="'+x.k+'" rows="2">'+esc(v||"")+'</textarea>':'<input type="text" data-ext="'+x.k+'" value="'+esc(v||"")+'">')+q}).join("");
  const link=D.artifact?(r.url||""):(r.pdf?'pdf/'+encodeURIComponent(r.pdf):"");
  R.innerHTML=
   '<div class="nav"><button id="prev">◀ 前 (←/k)</button><b>'+(cur+1)+' / '+view.length+'</b><button id="next">次 ▶ (→/j)</button>'
   +'<select id="flt">'+fltOpts()+'</select></div>'
   +'<h1>'+esc(r.title)+'</h1><div class="meta">'+esc(r.id)+' ・ '+esc(r.year)+' ・ '+esc(r.venue)+' ・ DOI '+esc(r.doi||"—")+'<br>担当ペアの相手: '+esc(r.partner)
   +' ・ 抽出担当: <b>'+(r.extractor==="Y"?"あなた":"相手")+'</b>'+(link?' ・ <a href="'+esc(link)+'" target="_blank" rel="noopener">'+(D.artifact?(r.oa?"論文を開く(OA)":"論文を開く(DOI)"):"PDF を別タブで開く")+'</a>':'')+'</div>'
   +'<div class="sec"><h2>① 基準の判定(Y / N / ?)</h2>'+crit
   +'<div style="margin-top:8px">判定(自動): <span id="decision"></span> <span class="meta" id="rn"></span></div></div>'
   +'<div class="sec"><h2>② 除外理由・根拠・メモ</h2>'
   +'<label class="f">除外理由(Exclude のとき自動。変えたいときだけ選び直す)</label><select id="reason"></select>'
   +'<label class="f">根拠(頁・引用) — Exclude では必須</label><textarea id="evidence">'+esc(e.evidence)+'</textarea>'
   +'<label class="f">メモ(迷った点・MR など)</label><textarea id="memo">'+esc(e.memo)+'</textarea></div>'
   +'<div class="sec" id="exsec"><h2>③ 事実の抽出 <span class="meta" id="exnote"></span></h2>'+exHtml+'</div>'
   +'<div id="probs"></div>'
   +'<div id="legend"><b>キー</b> ←/→ または k/j: 前後 ・ 1〜6: 基準を選ぶ ・ y / n / ?(または /): 付けて次の基準へ ・ 0: 消す ・ e: 根拠 ・ m: メモ ・ s: 保存 ・ l: 読込 ・ g: PDF 表示 ・ Tab: 一覧 ・ Esc: 入力を抜ける</div>';
  const opt=$("reason");opt.innerHTML='<option value="">(自動)</option>'+D.reasons.map(o=>'<option'+(e.reason===o?" selected":"")+'>'+esc(o)+'</option>').join("");
  bindFlt();
  $("pdfopen").style.display="";if(link){$("pdfopen").href=link}else{$("pdfopen").removeAttribute("href")}
  if(D.artifact){$("pdf").style.display="none";$("nopdf").style.display="";$("pdfopen").textContent=r.oa?"論文を開く(オープンアクセス)":"論文を開く(DOI・学外では読めないことがあります)";
    $("nopdf").innerHTML='<b>'+esc(r.title)+'</b><br>'+esc(r.venue)+' ('+esc(r.year)+')<br><br>上のボタンで別タブに開きます。読めないときは大学の VPN か、PC の pdf/ フォルダの「'+esc(r.pdf||"—")+'」を見てください。'}
  else if(pdfOn&&link){$("pdf").style.display="";$("nopdf").style.display="none";if($("pdf").dataset.src!==link){$("pdf").dataset.src=link;$("pdf").src=link}}
  else{$("pdf").style.display="none";$("nopdf").style.display="";$("nopdf").textContent=link?"(PDF 非表示。G で表示)":"この論文の PDF がありません。record_id と一緒に著者へ連絡してください。"}
  refresh();
}
function fltOpts(){return [["all","すべて"],["todo","判定が未了"],["Include","Include"],["Exclude","Exclude"],["Unsure","Unsure"],["ext","抽出が未了"],["prob","要修正あり"]]
  .map(([v,l])=>'<option value="'+v+'"'+(filter===v?" selected":"")+'>'+l+'</option>').join("")}
function bindFlt(){const f=$("flt");if(f)f.onchange=()=>{filter=f.value;cur=0;draw()}}

function refresh(){
  const r=D.records[view[cur]],e=E(r.id),d=dec(e);
  const dc=$("decision");dc.textContent=d||"(6 基準が未記入)";dc.className=d;
  $("rn").textContent=d==="Exclude"?("除外理由: "+(reasonOf(e)||"—")):"";
  const sel=$("reason");sel.disabled=d!=="Exclude";
  const ne=needExt(r,e);
  $("exsec").classList.toggle("dim",!ne);
  $("exnote").textContent=r.extractor!=="Y"?"(あなたは抽出担当ではありません。書かなくてよい)":(ne?"(Include/Unsure の論文: ● は必須。論文に無い事実は「記載なし」)":"(判定が Include か Unsure になったら書く)");
  const p=problems(r,e);
  $("probs").className=p.length?"":"okay";
  $("probs").innerHTML=!d?"":(p.length?"未記入・要修正: "+p.map(esc).join(" / "):"✔ この論文は記入できています");
  $("mTabF").textContent="✏ 記入 "+(cur+1)+"/"+view.length;
  document.querySelectorAll(".crit").forEach(el=>el.classList.toggle("on",+el.dataset.i===fc));
  const s=stats();
  $("stats").textContent="判定 "+s.done+"/"+s.n+"(Include "+s.inc+" ・ Exclude "+s.exc+" ・ Unsure "+s.uns+")  抽出 "+s.extDone+"/"+s.ext;
  $("bar").firstElementChild.style.width=(100*s.done/Math.max(s.n,1))+"%";
  drawSaved();
}
const SRV=/^https?:/.test(location.protocol);let srvTimer=null;
async function srvSave(){
  try{const r=await fetch("save",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(buildProgress())});
      if(!r.ok)throw new Error(r.status);dirty=false;LS.setItem("saved:"+KEY,new Date().toISOString());drawSaved();return true}
  catch(e){toast("PC への保存に失敗しました。通信を確認して、もう一度「進捗を保存」");return false}}
function changed(){
  if(!noSave)try{LS.setItem(KEY,JSON.stringify(store))}catch(e){}
  if(SRV&&!noSave){clearTimeout(srvTimer);srvTimer=setTimeout(srvSave,1200)}
  dirty=true;refresh();
  clearTimeout(saveTimer);saveTimer=setTimeout(()=>{if(handle)writeHandle()},1500);
}
function setCrit(k,v){const r=D.records[view[cur]],e=E(r.id);if(v==="")delete e.c[k];else e.c[k]=v;
  document.querySelectorAll('.seg[data-k="'+k+'"] button').forEach(b=>b.classList.toggle("on",b.dataset.v===v));changed()}
function go(n){const m=cur+n;if(m>=0&&m<view.length){cur=m;draw();$("right").scrollTop=0;window.scrollTo(0,0)}}

document.addEventListener("click",ev=>{
  const t=ev.target;
  if(t.closest(".seg[data-k]")&&t.tagName==="BUTTON"){const k=t.parentElement.dataset.k,e=E(D.records[view[cur]].id);setCrit(k,e.c[k]===t.dataset.v?"":t.dataset.v);return}
  if(t.closest(".seg[data-ek]")&&t.tagName==="BUTTON"){const k=t.parentElement.dataset.ek,e=E(D.records[view[cur]].id);
    if(e.ext[k]===t.dataset.v)delete e.ext[k];else e.ext[k]=t.dataset.v;
    t.parentElement.querySelectorAll("button").forEach(b=>b.classList.toggle("on",e.ext[k]===b.dataset.v));changed();return}
  const c=t.closest(".crit");if(c&&!t.closest("button")){fc=+c.dataset.i;refresh();return}
  if(t.dataset&&t.dataset.fill){const e=E(D.records[view[cur]].id);e.ext[t.dataset.fill]=t.dataset.v;
    const inp=document.querySelector('[data-ext="'+t.dataset.fill+'"]');if(inp)inp.value=t.dataset.v;changed();return}
  if(t.id==="prev"||t.id==="mPrev")go(-1);if(t.id==="next"||t.id==="mNext")go(1);
  if(t.id==="mTabF")setTab("form");if(t.id==="mTabP")setTab("pdf");
  const row=t.closest("tr.row");if(row){filter="all";cur=0;rebuild();cur=view.findIndex(i=>D.records[i].id===row.dataset.id);$("list").classList.remove("on");draw()}
});
document.addEventListener("input",ev=>{
  const t=ev.target,e=view.length?E(D.records[view[cur]].id):null;if(!e)return;
  if(t.id==="evidence"){e.evidence=t.value;changed()}
  else if(t.id==="memo"){e.memo=t.value;changed()}
  else if(t.id==="reason"){e.reason=t.value;changed()}
  else if(t.dataset&&t.dataset.ext){if(t.value==="")delete e.ext[t.dataset.ext];else e.ext[t.dataset.ext]=t.value;changed()}
  else if(t.type==="checkbox"){const k=t.closest("[data-ms]").dataset.ms;
    e.ext[k]=[...t.closest("[data-ms]").querySelectorAll("input:checked")].map(i=>i.value);if(!e.ext[k].length)delete e.ext[k];changed()}
});
document.addEventListener("keydown",ev=>{
  if(ev.ctrlKey||ev.metaKey||ev.altKey)return;
  const t=ev.target,typing=t&&(t.tagName==="TEXTAREA"||(t.tagName==="INPUT"&&t.type==="text")||t.tagName==="SELECT");
  if(ev.key==="Escape"){if(typing)t.blur();$("list").classList.remove("on");return}
  if(typing)return;
  if($("list").classList.contains("on")&&ev.key!=="Tab")return;
  const k=ev.key;
  if(k==="Tab"){ev.preventDefault();$("list").classList.contains("on")?$("list").classList.remove("on"):showList();return}
  if(!view.length)return;
  if(k==="ArrowRight"||k==="j"){go(1)}else if(k==="ArrowLeft"||k==="k"){go(-1)}
  else if(/^[1-6]$/.test(k)){fc=+k-1;refresh()}
  else if(k==="y"||k==="Y"||k==="n"||k==="N"||k==="?"||k==="/"||k==="q"){
    const v=(k==="y"||k==="Y")?"Y":(k==="n"||k==="N")?"N":"?";setCrit(CR[fc].k,v);if(fc<CR.length-1){fc++;refresh()}}
  else if(k==="0"||k==="Backspace"){setCrit(CR[fc].k,"")}
  else if(k==="e"){ev.preventDefault();$("evidence").focus()}
  else if(k==="m"){ev.preventDefault();$("memo").focus()}
  else if(k==="s"){saveJSON()}else if(k==="l"){D.artifact?sheetBox("load"):$("file").click()}
  else if(k==="g"){togglePdf()}
});
function setTab(v){document.body.classList.toggle("m-pdf",v==="pdf");document.body.classList.toggle("m-form",v!=="pdf");
  $("mTabF").classList.toggle("on",v!=="pdf");$("mTabP").classList.toggle("on",v==="pdf")}
function togglePdf(){pdfOn=!pdfOn;$("main").classList.toggle("hide-pdf",!pdfOn);draw()}
$("bPdf").onclick=togglePdf;$("bList").onclick=()=>$("list").classList.contains("on")?$("list").classList.remove("on"):showList();
function showList(){
  const rows=D.records.map((r,i)=>{const e=store[r.id]||{c:{},ext:{}},d=dec(e),p=problems(r,E(r.id));
    return '<tr class="row" data-id="'+esc(r.id)+'"><td>'+(i+1)+'</td><td>'+esc(r.id)+'</td><td>'+esc(r.title.slice(0,90))+'</td><td>'+(r.extractor==="Y"?"抽出":"")+'</td><td>'+esc(d||"—")+'</td><td style="color:#b42318">'+(d&&p.length?"要修正 "+p.length:"")+'</td></tr>'}).join("");
  $("list").innerHTML='<table><tr><th>#</th><th>ID</th><th>タイトル</th><th>担当</th><th>判定</th><th></th></tr>'+rows+'</table>';$("list").classList.add("on")}

// ---- 保存(進捗 JSON) --------------------------------------------------------------------
function buildProgress(){
  const entries={};for(const r of D.records){const e=store[r.id];if(e&&touched(e))entries[r.id]=e}
  const s=stats();
  return {format:"phase4-progress",version:1,sheet:D.sheet,fingerprint:D.fingerprint,saved_at:new Date().toISOString(),
          counts:{done:s.done,include:s.inc,exclude:s.exc,unsure:s.uns,extraction_done:s.extDone,extraction_needed:s.ext,total:s.n},entries}}
function toast(m){const t=$("toast");t.textContent=m;t.style.display="block";clearTimeout(toast.t);toast.t=setTimeout(()=>t.style.display="none",3500)}
async function writeHandle(){
  try{const w=await handle.createWritable();await w.write(JSON.stringify(buildProgress(),null,1));await w.close();
      dirty=false;LS.setItem("saved:"+KEY,new Date().toISOString());drawSaved()}catch(e){handle=null;toast("自動保存に失敗しました。S で保存し直してください")}}
function sheetBox(mode){
  const box=$("jbox");box.style.display="flex";const ta=$("jtext");
  if(mode==="save"){ta.value=JSON.stringify(buildProgress());ta.readOnly=true;$("jtitle").textContent="進捗 JSON(コピーして PC の progress_"+D.sheet+".json に保存するか、Claude に貼る)";$("jgo").textContent="コピー"}
  else{ta.value="";ta.readOnly=false;$("jtitle").textContent="進捗 JSON を貼り付けて「読み込む」";$("jgo").textContent="読み込む"}
  box.dataset.mode=mode}
async function jgo(){
  const ta=$("jtext");
  if($("jbox").dataset.mode==="save"){
    try{await navigator.clipboard.writeText(ta.value)}catch(e){ta.select();try{document.execCommand("copy")}catch(_){}}
    dirty=false;LS.setItem("saved:"+KEY,new Date().toISOString());drawSaved();toast("コピーしました");return}
  loadProgress(ta.value)}
async function saveJSON(){
  if(D.artifact){sheetBox("save");return}
  if(SRV){if(await srvSave())toast("PC に保存しました(以後、変更のたびに自動で保存します)");return}
  if(!handle&&window.showSaveFilePicker){
    try{handle=await window.showSaveFilePicker({suggestedName:"progress_"+D.sheet+".json",types:[{description:"進捗 JSON",accept:{"application/json":[".json"]}}]});}
    catch(e){if(e&&e.name==="AbortError")return;handle=null}}
  if(handle){await writeHandle();toast("保存しました(以後、変更のたびにこのファイルへ自動保存します)");return}
  const a=document.createElement("a");a.href=URL.createObjectURL(new Blob([JSON.stringify(buildProgress(),null,1)],{type:"application/json"}));
  a.download="progress_"+D.sheet+".json";document.body.appendChild(a);a.click();a.remove();
  dirty=false;LS.setItem("saved:"+KEY,new Date().toISOString());drawSaved();toast("ダウンロードしました。このファイルを screening/phase4/"+D.sheet+"/ に置いてください")}
function drawSaved(){
  const t=LS.getItem("saved:"+KEY);
  $("saved").textContent=(dirty?"● 未保存 ・ ":"")+(t?"最終保存 "+t.slice(5,16).replace("T"," "):"JSON は未保存(S で保存)");
}
$("bSave").onclick=saveJSON;$("bLoad").onclick=()=>D.artifact?sheetBox("load"):$("file").click();
$("jgo").onclick=jgo;$("jclose").onclick=()=>$("jbox").style.display="none";
$("file").onchange=async()=>{
  const f=$("file").files[0];if(!f)return;$("file").value="";loadProgress(await f.text())};
function loadProgress(text){
  let p;try{p=JSON.parse(text)}catch(e){toast("JSON として読めません");return}
  if(p.format!=="phase4-progress"){toast("このアプリの進捗 JSON ではありません");return}
  if(p.sheet!==D.sheet){toast("他の評価者の進捗 JSON です。読み込めません");return}
  if(p.fingerprint!==D.fingerprint){if(!confirm("担当の論文一覧が、この HTML と違います(作り直した可能性)。それでも読み込みますか?"))return}
  const n=Object.keys(p.entries||{}).length;
  if(!confirm("進捗 JSON(保存 "+(p.saved_at||"?").slice(0,16).replace("T"," ")+"、"+n+" 件)を読み込みます。\nいまのブラウザ内の入力は置き換わります。よいですか?"))return;
  LS.setItem(KEY+":backup",JSON.stringify(store));
  store={};for(const [id,e] of Object.entries(p.entries||{}))store[id]=e;
  changed();draw();$("jbox").style.display="none";toast("読み込みました("+n+" 件)")}
window.addEventListener("beforeunload",e=>{if(dirty){e.preventDefault();e.returnValue=""}});

// ---- 自己テスト(`…html#selftest` で開くと、判定の規則が式と同じか確認して画面に書く) -------------
function selftest(){
  const out=[],r=D.records[0],had=store[r.id];const bak=JSON.stringify(had||null);
  const e={c:{},reason:"",evidence:"",memo:"",ext:{}};store[r.id]=e;
  const all=v=>{for(const x of CR)e.c[x.k]=v};
  all("Y");out.push(["全 Y → Include",dec(e)==="Include"]);
  e.c.C="N";out.push(["C だけ N → Exclude / その他",dec(e)==="Exclude"&&autoR(e)===D.other]);
  all("Y");e.c.S="N";e.c.P="N";out.push(["S と P が N → S の理由",autoR(e)==="S: ユーザー実験が無い"]);
  all("Y");e.c.O="N";out.push(["O が N",autoR(e)==="O: スケール知覚の測定が無い"]);
  all("Y");e.c.I_hmd="?";out.push(["? → Unsure(理由なし)",dec(e)==="Unsure"&&reasonOf(e)===""]);
  delete e.c.S;all("Y");delete e.c.S;out.push(["空欄 → 判定なし",dec(e)===""]);
  all("N");e.reason="重複(同一研究の別報告)";out.push(["選び直した理由が優先",reasonOf(e)==="重複(同一研究の別報告)"]);
  e.evidence="";out.push(["Exclude+根拠なし → 要修正",problems(r,e).includes("根拠(頁・引用)")]);
  all("Y");e.ext.n_total="abc";out.push(["n_total が非整数 → 要修正",problems(r,e).some(x=>x.includes("整数"))]);
  e.ext.n_total="24";out.push(["n_total=24 は可",!problems(r,e).some(x=>x.includes("整数"))]);
  if(had)store[r.id]=JSON.parse(bak);else delete store[r.id];
  // キー操作: 1 件目で y,y,y,y,y,y → Include、続けて 4 → C を選んで n → Exclude(その他)
  noSave=true;const first=D.records[0].id,hadFirst=JSON.stringify(store[first]||null);delete store[first];filter="all";cur=0;fc=0;draw();
  for(const k of["y","y","y","y","y","y"])document.dispatchEvent(new KeyboardEvent("keydown",{key:k}));
  out.push(["キー y×6 → Include",dec(E(first))==="Include"&&$("decision").textContent==="Include"]);
  document.dispatchEvent(new KeyboardEvent("keydown",{key:"4"}));document.dispatchEvent(new KeyboardEvent("keydown",{key:"n"}));
  out.push(["キー 4,n → Exclude/その他",dec(E(first))==="Exclude"&&$("rn").textContent.includes(D.other)]);
  document.dispatchEvent(new KeyboardEvent("keydown",{key:"ArrowRight"}));
  out.push(["→ で次の論文",D.records[view[cur]].id===D.records[1].id]);
  if(hadFirst!=="null")store[first]=JSON.parse(hadFirst);else delete store[first];noSave=false;
  const bad=out.filter(x=>!x[1]);
  $("selftest").textContent=bad.length?"SELFTEST FAIL "+JSON.stringify(bad):"SELFTEST OK "+out.length;
}
setTab("form");
draw();
if(SRV&&location.hash!=="#selftest"){      // サーバー経由なら、PC 側の保存を引き継ぐ(スマホのブラウザは空なので)
  fetch("progress.json").then(r=>r.ok?r.json():null).then(p=>{
    if(!p||p.format!=="phase4-progress"||p.sheet!==D.sheet)return;
    if(D.records.some(r=>store[r.id]&&touched(store[r.id])&&!r.init))return;    // 既にこのブラウザに入力がある
    for(const [id,e] of Object.entries(p.entries||{}))if(!store[id]||!touched(store[id]))store[id]=e;
    draw();toast("PC に保存済みの進捗を読み込みました")}).catch(()=>{})}
if(location.hash==="#selftest")selftest();
</script></body></html>
"""

if __name__ == "__main__":
    main()
